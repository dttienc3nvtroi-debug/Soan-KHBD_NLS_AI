import io
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import google.generativeai as genai
import streamlit as st

# Cấu hình giao diện Streamlit
st.set_page_config(
    page_title="Hệ Thống Soạn Giáo Án AI - GDPT 2018",
    page_icon="📚",
    layout="wide",
)

st.title("📚 HỆ THỐNG SOẠN GIÁO ÁN TÍCH HỢP AI & NĂNG LỰC SỐ")
st.caption(
    "Chuẩn Công văn 5512/2345 & Tích hợp Giáo dục AI (QĐ 2422/QĐ-BGDĐT & CV"
    " 5588/BGDĐT-GDPT)"
)

# Lấy API Key từ Secrets
raw_api_key = st.secrets.get("GEMINI_API_KEY", "")
api_key = str(raw_api_key).strip().strip('"').strip("'")

# GIAO DIỆN NHẬP THÔNG TIN
st.subheader("1. Thông tin chung bài dạy")
col1, col2, col3 = st.columns(3)
with col1:
    truong = st.text_input("Trường THPT/THCS:", "THPT Nguyễn Văn Trỗi")
    gv = st.text_input("Họ và tên giáo viên:", "Dương Tấn Tiến")
    mon = st.text_input("Môn học:", "Toán học")
with col2:
    to = st.text_input("Tổ chuyên môn:", "Tổ Toán")
    lop = st.selectbox(
        "Khối lớp:",
        ["Lớp 10", "Lớp 11", "Lớp 12", "Lớp 6", "Lớp 7", "Lớp 8", "Lớp 9"],
    )
    thoi_luong = st.text_input("Thời lượng thực hiện:", "2 tiết")
with col3:
    sach = st.selectbox(
        "Bộ sách giáo khoa:",
        ["Kết nối tri thức với cuộc sống", "Cánh diều", "Chân trời sáng tạo"],
    )
    ten_bai = st.text_input(
        "Tên bài dạy:",
        "BÀI 5: GIÁ TRỊ LƯỢNG GIÁC CỦA MỘT GÓC TỪ 0° ĐẾN 180°",
    )
    ngay_soan = st.date_input("Ngày soạn:")

st.subheader("2. Đặc điểm lớp học & Phân hóa học sinh")
col_tl1, col_tl2 = st.columns([1, 2])
with col_tl1:
    so_hs = st.number_input("Tổng số học sinh:", value=40)
    hs_khuyet_tat = st.checkbox("Có học sinh khuyết tật hòa nhập", value=True)
with col_tl2:
    st.write("Tỉ lệ phân hóa học lực học sinh:")
    gioi = st.slider("Học sinh Giỏi (%):", 0, 100, 15)
    kha = st.slider("Học sinh Khá (%):", 0, 100, 55)
    tb = st.slider("Học sinh Trung bình (%):", 0, 100, 25)
    yeu = 100 - (gioi + kha + tb)
    st.info(f"👉 Tỉ lệ HS Yếu/Kém tính toán tự động: **{yeu}%**")

st.subheader("3. Yêu cầu Cần Đạt / Mục tiêu bài học (Chính xác tuyệt đối)")
use_custom_objectives = st.checkbox(
    "📌 Cung cấp Mục tiêu / Yêu cầu cần đạt chuẩn (Kiến thức, Năng lực, Phẩm"
    " chất)",
    value=True,
)

custom_objectives_text = ""
if use_custom_objectives:
    custom_objectives_text = st.text_area(
        "Dán/Nhập trực tiếp Mục tiêu bài học vào đây:",
        height=180,
        value="""1. Kiến thức:
- Nhận biết và định nghĩa được giá trị lượng giác (sin, cos, tan, cot) của một góc từ 0° đến 180° trên nửa đường tròn đơn vị.
- Giải thích được mối quan hệ giữa các giá trị lượng giác của hai góc bù nhau.
- Nắm được bảng giá trị lượng giác góc đặc biệt và biết dùng MTCT.

2. Năng lực & Phẩm chất:
- Năng lực toán học: Tư duy, luận lý và giải quyết vấn đề toán học.
- Năng lực Số/AI: Tương tác qua GeoGebra, Kahoot, tra cứu Google Maps và ứng dụng AI (ChatGPT/Canva) hỗ trợ học tập.
- Phẩm chất: Trách nhiệm (hỗ trợ bạn khuyết tật hòa nhập), chăm chỉ, trung thực.""",
    )

st.subheader("4. Tích hợp Công nghệ & AI")
col_tc1, col_tc2 = st.columns(2)
with col_tc1:
    cn_su_dung = st.multiselect(
        "Công nghệ/Nền tảng số sử dụng:",
        ["GeoGebra", "Kahoot", "Padlet", "Google Maps", "Canva AI"],
        default=["GeoGebra", "Kahoot", "Padlet", "Google Maps"],
    )
with col_tc2:
    ai_tich_hop = st.multiselect(
        "Tích hợp Giáo dục Trí tuệ Nhân tạo (AI):",
        ["ChatGPT/Gemini làm Trợ lý học tập", "Canva AI thiết kế Infographic"],
        default=[
            "ChatGPT/Gemini làm Trợ lý học tập",
            "Canva AI thiết kế Infographic",
        ],
    )

# NÚT BẤM KÍCH HOẠT SOẠN GIÁO ÁN
if st.button("🚀 SOẠN GIÁO ÁN TÍCH HỢP TỰ ĐỘNG", type="primary"):
    if not api_key:
        st.error(
            "Chưa cấu hình GEMINI_API_KEY trong phần Secrets trên Streamlit!"
        )
    else:
        with st.spinner("AI đang soạn thảo Kế hoạch bài dạy chi tiết..."):
            try:
                genai.configure(api_key=api_key)

                prompt = f"""
Hãy đóng vai là Giáo viên giỏi môn {mon}. Soạn Kế hoạch bài dạy (Giáo án) theo CV 5512, tích hợp Giáo dục AI (QĐ 2422 & CV 5588).

THÔNG TIN BÀI DẠY:
- Trường: {truong} | Tổ: {to} | GV: {gv}
- Tên bài: {ten_bai} | Lớp: {lop} | Thời lượng: {thoi_luong} | Bộ sách: {sach}
- Đặc điểm lớp: {so_hs} HS ({gioi}% Giỏi, {kha}% Khá, {tb}% TB, {yeu}% Yếu. {'Có 1 HS khuyết tật hòa nhập' if hs_khuyet_tat else ''}).
- Công nghệ số: {', '.join(cn_su_dung)}
- Tích hợp AI: {', '.join(ai_tich_hop)}

MỤC TIÊU BÀI HỌC (YÊU CẦU CẦN ĐẠT CUNG CẤP CHÍNH XÁC):
{custom_objectives_text if use_custom_objectives else 'Tự đề xuất mục tiêu chuẩn theo GDPT 2018'}

YÊU CẦU CẤU TRÚC GIÁO ÁN:
I. MỤC TIÊU (Giữ chuẩn mục tiêu đã cung cấp, bổ sung chi tiết Năng lực đặc thù, Năng lực chung, Năng lực số/AI và Phẩm chất)
II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU (GV & HS)
III. TIẾN TRÌNH DẠY HỌC (Đủ 4 hoạt động: Mở đầu, Hình thành kiến thức, Luyện tập - bám sát SGK, Vận dụng tích hợp AI/Google Maps). Mỗi hoạt động trình bày rõ ràng 4 bước: Bước 1 Chuyển giao, Bước 2 Thực hiện, Bước 3 Báo cáo, Bước 4 Kết luận.
""".strip()

                model = genai.GenerativeModel("gemini-1.5-flash")
                response = model.generate_content(prompt)
                plan_text = response.text

                st.success(" Soạn giáo án hoàn tất!")
                st.markdown(plan_text)

                # TẠO FILE WORD ĐỂ TẢI VỀ
                doc = docx.Document()
                table = doc.add_table(rows=1, cols=2)
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = f"TRƯỜNG: {truong.upper()}\nTỔ: {to.upper()}"
                hdr_cells[1].text = (
                    f"NGÀY SOẠN: {ngay_soan.strftime('%d/%m/%Y')}\nGV:"
                    f" {gv.upper()}"
                )

                p_title = doc.add_paragraph()
                p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_title = p_title.add_run(
                    f"\nTÊN BÀI DẠY: {ten_bai.upper()}\n"
                )
                run_title.bold = True
                run_title.font.size = Pt(14)

                p_sub = doc.add_paragraph(
                    f"Môn học: {mon}; Lớp: {lop} ({thoi_luong})\n"
                )
                p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

                doc.add_paragraph(plan_text)

                bio = io.BytesIO()
                doc.save(bio)

                st.download_button(
                    label="📥 TẢI FILE WORD (.DOCX) VỀ MÁY TÍNH",
                    data=bio.getvalue(),
                    file_name=f"GiaoAn_{ten_bai.replace(' ', '_')}.docx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    ),
                    type="primary",
                )
            except Exception as e:
                st.error(
                    f"Có lỗi xảy ra khi kết nối với Gemini API. Chi tiết:"
                    f" {str(e)}"
                )
                st.info(
                    "💡 Hướng dẫn sửa: Hãy kiểm tra lại phần Secrets trên"
                    " Streamlit Cloud. Đảm bảo GEMINI_API_KEY bắt đầu bằng"
                    " 'AIzaSy...' được lấy từ Google AI Studio"
                    " (aistudio.google.com)."
                )
